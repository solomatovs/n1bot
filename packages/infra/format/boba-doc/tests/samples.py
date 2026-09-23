"""Образцы документов для тестов: настоящие файлы, собранные библиотеками
форматов или руками там, где библиотека только читает."""

from __future__ import annotations

import io
import os
import threading
from collections.abc import Generator, Mapping, Sequence
from contextlib import contextmanager
from pathlib import Path
from typing import Any, BinaryIO

import docx
import openpyxl
import xlwt
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches

from boba.doc.document import ByteStream


class NoSeek(ByteStream):
    """Поток только с read: у него нет ни seek, ни tell, как у сокета."""

    def __init__(self, data: bytes) -> None:
        self._buffer = io.BytesIO(data)

    def read(self, size: int = -1, /) -> bytes:
        return self._buffer.read(size)


class PipeSource:
    """Настоящий файловый дескриптор: байты приходят по пипе из другого потока."""

    @staticmethod
    @contextmanager
    def of(data: bytes) -> Generator[BinaryIO, None, None]:
        read_end, write_end = os.pipe()

        def feed() -> None:
            with os.fdopen(write_end, "wb") as sink:
                sink.write(data)

        writer = threading.Thread(target=feed)
        writer.start()
        try:
            with os.fdopen(read_end, "rb") as source:
                yield source
        finally:
            writer.join()


class Samples:
    """Сборка образцов по форматам."""

    FONT_SIZE = 28
    LINE_HEIGHT = 48

    @staticmethod
    def pdf(pages: Sequence[str]) -> bytes:
        """Многостраничный pdf с текстовым слоем; смещения xref честные."""
        objects: list[bytes] = [b"", b""]
        font_number = 3
        objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

        kids: list[bytes] = []
        for text in pages:
            stream = f"BT /F1 18 Tf 40 700 Td ({text}) Tj ET".encode("latin-1")
            page_number = len(objects) + 1
            content_number = page_number + 1
            kids.append(f"{page_number} 0 R".encode())
            objects.append(
                b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                b"/Resources << /Font << /F1 "
                + f"{font_number} 0 R".encode()
                + b" >> >> /Contents "
                + f"{content_number} 0 R".encode()
                + b" >>"
            )
            objects.append(
                b"<< /Length "
                + str(len(stream)).encode()
                + b" >>\nstream\n"
                + stream
                + b"\nendstream"
            )

        objects[0] = b"<< /Type /Catalog /Pages 2 0 R >>"
        objects[1] = (
            b"<< /Type /Pages /Kids ["
            + b" ".join(kids)
            + f"] /Count {len(pages)} >>".encode()
        )

        out = bytearray(b"%PDF-1.4\n")
        offsets: list[int] = []
        for number, body in enumerate(objects, start=1):
            offsets.append(len(out))
            out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"

        xref = len(out)
        out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
        for offset in offsets:
            out += f"{offset:010d} 00000 n \n".encode()

        out += (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref}\n%%EOF\n"
        ).encode()

        return bytes(out)

    @classmethod
    def image(cls, lines: Sequence[str], font: Path) -> Image.Image:
        typeface = ImageFont.truetype(str(font), cls.FONT_SIZE)
        height = cls.LINE_HEIGHT * (len(lines) + 1)
        image = Image.new("RGB", (1000, height), "white")
        draw = ImageDraw.Draw(image)
        for index, line in enumerate(lines):
            draw.text(
                (20, 20 + index * cls.LINE_HEIGHT), line, font=typeface, fill="black"
            )

        return image

    @classmethod
    def png(cls, lines: Sequence[str], font: Path) -> bytes:
        buffer = io.BytesIO()
        cls.image(lines, font).save(buffer, format="PNG")

        return buffer.getvalue()

    @classmethod
    def tiff(cls, frames: Sequence[Sequence[str]], font: Path) -> bytes:
        images = [cls.image(lines, font) for lines in frames]
        buffer = io.BytesIO()
        images[0].save(buffer, format="TIFF", save_all=True, append_images=images[1:])

        return buffer.getvalue()

    @classmethod
    def scanned_pdf(cls, lines: Sequence[str], font: Path) -> bytes:
        """PDF из картинки без текстового слоя, как со сканера."""
        buffer = io.BytesIO()
        cls.image(lines, font).save(buffer, format="PDF")

        return buffer.getvalue()

    @staticmethod
    def docx(paragraphs: Sequence[str], table: Sequence[Sequence[str]]) -> bytes:
        document = docx.Document()
        for text in paragraphs:
            document.add_paragraph(text)

        if table:
            grid = document.add_table(rows=len(table), cols=len(table[0]))
            for row_index, row in enumerate(table):
                for column_index, value in enumerate(row):
                    grid.cell(row_index, column_index).text = value

        buffer = io.BytesIO()
        document.save(buffer)

        return buffer.getvalue()

    @staticmethod
    def xlsx(sheets: Mapping[str, Sequence[Sequence[object]]]) -> bytes:
        workbook = openpyxl.Workbook()
        default = workbook.active
        if default is not None:
            workbook.remove(default)

        for title, rows in sheets.items():
            sheet = workbook.create_sheet(title)
            for row in rows:
                sheet.append(list(row))

        buffer = io.BytesIO()
        workbook.save(buffer)

        return buffer.getvalue()

    @staticmethod
    def xls(sheets: Mapping[str, Sequence[Sequence[object]]]) -> bytes:
        workbook = xlwt.Workbook()
        for title, rows in sheets.items():
            sheet = workbook.add_sheet(title)
            for row_index, row in enumerate(rows):
                for column_index, value in enumerate(row):
                    if value is None:
                        continue

                    sheet.write(row_index, column_index, value)

        buffer = io.BytesIO()
        workbook.save(buffer)

        return buffer.getvalue()

    @staticmethod
    def pptx(
        slides: Sequence[str], table: Sequence[Sequence[str]], notes: str
    ) -> bytes:
        presentation = Presentation()
        layout = presentation.slide_layouts[5]
        for text in slides:
            slide = presentation.slides.add_slide(layout)
            Samples._set_title(slide, text)

        if table:
            slide = presentation.slides.add_slide(layout)
            Samples._set_title(slide, "Table slide")
            shape = slide.shapes.add_table(
                len(table), len(table[0]), Inches(1), Inches(2), Inches(6), Inches(2)
            )
            for row_index, row in enumerate(table):
                for column_index, value in enumerate(row):
                    shape.table.cell(row_index, column_index).text = value

            frame = slide.notes_slide.notes_text_frame
            if frame is not None:
                frame.text = notes

        buffer = io.BytesIO()
        presentation.save(buffer)

        return buffer.getvalue()

    @staticmethod
    def _set_title(slide: Any, text: str) -> None:
        title = slide.shapes.title
        if title is None:
            raise AssertionError("slide layout without a title placeholder")

        title.text = text

    @staticmethod
    def rtf(text: str) -> bytes:
        """RTF в cp1251: не-ASCII символы записаны hex-эскейпами."""
        escaped: list[str] = []
        for char in text:
            if ord(char) < 128:
                escaped.append(char)
                continue

            escaped.append("\\'" + char.encode("cp1251").hex())

        body = "".join(escaped)
        markup = (
            "{\\rtf1\\ansi\\ansicpg1251\\deff0{\\fonttbl{\\f0 Arial;}}\\f0\\fs24 "
            + body
            + "\\par}"
        )

        return markup.encode("ascii")
